import logging
import os
from io import BytesIO

from openpyxl import load_workbook

from fastapi import APIRouter
from fastapi import Depends
from fastapi import File
from fastapi import Form
from fastapi import HTTPException
from fastapi import UploadFile
from fastapi.responses import FileResponse

from sqlalchemy.orm import Session

from application_state import ROUND_STAGES
from application_state import advance_application_for_round
from application_state import is_application_ready_for_round
from application_state import next_stage_after_round
from application_state import normalize_stage
from application_state import reject_application_for_round
from database import get_db

from models import Application
from models import PlacementDrive
from models import Student
from models import User

from schemas import DriveCreate
from schemas import DriveResponse
from schemas import DriveUpdate
from upload_security import EXCEL_MIME_TYPES
from upload_security import JD_MIME_TYPES
from upload_security import MAX_EXCEL_SIZE
from upload_security import MAX_JD_SIZE
from upload_security import generate_safe_upload_path
from upload_security import read_upload_with_limit
from upload_security import stored_upload_path
from upload_security import validate_excel_content
from upload_security import validate_extension_and_mime
from upload_security import validate_pdf_content
from upload_security import write_upload_with_limit

from routers.auth import get_current_user
from routers.auth import require_admin


logger = logging.getLogger(__name__)


router = APIRouter(
    prefix="/api/drives",
    tags=["Placement Drives"],
)


UPLOAD_DIR = "uploads/jd"

os.makedirs(
    UPLOAD_DIR,
    exist_ok=True
)


def normalize_roll_no(value):
    if value is None:
        return ""

    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("\u00a0", "")
    )


@router.get(
    "/",
    response_model=list[DriveResponse],
)
def get_drives(
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(
        get_db
    ),
):
    user_role = current_user.role.lower()

    if user_role == "admin":
        query = db.query(PlacementDrive)
    elif user_role == "student":
        query = db.query(PlacementDrive).filter(
            PlacementDrive.status == "Published"
        )
    else:
        raise HTTPException(
            status_code=403,
            detail="You are not authorized to view placement drives",
        )

    return (
        query
        .order_by(
            PlacementDrive.id.desc()
        )
        .all()
    )


@router.get(
    "/{drive_id}",
    response_model=DriveResponse,
)
def get_drive(
    drive_id: int,
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(
        get_db
    ),
):
    drive = (
        db.query(
            PlacementDrive
        )
        .filter(
            PlacementDrive.id
            == drive_id
        )
        .first()
    )

    if not drive:
        raise HTTPException(
            status_code=404,
            detail=(
                "Placement drive not found"
            ),
        )

    if (
        current_user.role.lower() == "student"
        and drive.status != "Published"
    ):
        raise HTTPException(
            status_code=404,
            detail="Placement drive not found",
        )

    if current_user.role.lower() != "admin" and current_user.role.lower() != "student":
        raise HTTPException(
            status_code=403,
            detail="You are not authorized to view placement drives",
        )

    return drive


@router.post(
    "/",
    response_model=DriveResponse,
)
def create_drive(
    drive_data: DriveCreate,
    current_user: User = Depends(
        require_admin
    ),
    db: Session = Depends(
        get_db
    ),
):
    try:
        drive = PlacementDrive(
            **drive_data.model_dump()
        )

        db.add(drive)

        db.commit()

        db.refresh(drive)

        return drive

    except Exception as error:
        db.rollback()

        logger.exception(
            "Unexpected error while creating placement drive"
        )

        raise HTTPException(
            status_code=500,
            detail="Failed to create placement drive.",
        ) from error


@router.patch(
    "/{drive_id}",
    response_model=DriveResponse,
)
def update_drive(
    drive_id: int,
    drive_data: DriveUpdate,
    current_user: User = Depends(
        require_admin
    ),
    db: Session = Depends(
        get_db
    ),
):
    drive = (
        db.query(
            PlacementDrive
        )
        .filter(
            PlacementDrive.id
            == drive_id
        )
        .first()
    )

    if not drive:
        raise HTTPException(
            status_code=404,
            detail=(
                "Placement drive not found"
            ),
        )

    update_data = (
        drive_data.model_dump(
            exclude_unset=True
        )
    )

    try:
        for field, value in (
            update_data.items()
        ):
            setattr(
                drive,
                field,
                value,
            )

        db.commit()

        db.refresh(drive)

        return drive

    except Exception as error:
        db.rollback()

        logger.exception(
            "Unexpected error while updating placement drive"
        )

        raise HTTPException(
            status_code=500,
            detail="Failed to update placement drive.",
        ) from error


@router.post(
    "/{drive_id}/jd",
    response_model=DriveResponse,
)
async def upload_job_description(
    drive_id: int,
    file: UploadFile = File(...),
    current_user: User = Depends(
        require_admin
    ),
    db: Session = Depends(
        get_db
    ),
):
    drive = (
        db.query(
            PlacementDrive
        )
        .filter(
            PlacementDrive.id
            == drive_id
        )
        .first()
    )

    if not drive:
        raise HTTPException(
            status_code=404,
            detail=(
                "Placement drive not found"
            ),
        )

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected",
        )

    destination = None
    old_jd_path = stored_upload_path(
        UPLOAD_DIR,
        drive.jd,
        "/uploads/jd",
    )

    try:
        extension = validate_extension_and_mime(
            file,
            JD_MIME_TYPES,
            "Only PDF files are allowed",
        )

        destination = generate_safe_upload_path(
            UPLOAD_DIR,
            extension,
        )

        await write_upload_with_limit(
            file,
            destination,
            MAX_JD_SIZE,
        )

        validate_pdf_content(
            destination
        )

        drive.jd = (
            f"/uploads/jd/{destination.name}"
        )

        drive.jd_filename = (
            os.path.basename(
                file.filename.replace(
                    "\\",
                    "/",
                )
            )
        )

        db.commit()

        db.refresh(drive)

        if (
            old_jd_path
            and old_jd_path.exists()
            and old_jd_path != destination
        ):
            try:
                old_jd_path.unlink()
            except OSError:
                logger.warning(
                    "Unable to remove replaced JD file",
                    exc_info=True,
                )

        return drive

    except HTTPException:
        db.rollback()

        if (
            destination
            and destination.exists()
        ):
            try:
                destination.unlink()
            except OSError:
                logger.warning(
                    "Unable to remove rejected JD upload",
                    exc_info=True,
                )

        raise

    except Exception as error:
        db.rollback()

        logger.exception(
            "Unexpected error while uploading job description"
        )

        if (
            destination
            and destination.exists()
        ):
            try:
                destination.unlink()
            except OSError:
                logger.warning(
                    "Unable to remove failed JD upload",
                    exc_info=True,
                )

        raise HTTPException(
            status_code=500,
            detail="Failed to upload job description.",
        ) from error

    finally:
        await file.close()


@router.get(
    "/{drive_id}/jd",
)
def download_job_description(
    drive_id: int,
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(get_db),
):
    drive = (
        db.query(PlacementDrive)
        .filter(
            PlacementDrive.id == drive_id
        )
        .first()
    )

    if not drive:
        raise HTTPException(
            status_code=404,
            detail="Placement drive not found",
        )

    user_role = current_user.role.lower()

    if (
        user_role == "student"
        and drive.status != "Published"
    ):
        raise HTTPException(
            status_code=404,
            detail="Placement drive not found",
        )

    if user_role not in {
        "admin",
        "student",
    }:
        raise HTTPException(
            status_code=403,
            detail=(
                "You are not authorized to "
                "access job descriptions"
            ),
        )

    jd_path = stored_upload_path(
        UPLOAD_DIR,
        drive.jd,
        "/uploads/jd",
    )

    if (
        not jd_path
        or not jd_path.is_file()
    ):
        raise HTTPException(
            status_code=404,
            detail="Job description not found",
        )

    download_name = os.path.basename(
        (
            drive.jd_filename
            or jd_path.name
        ).replace("\\", "/")
    )

    return FileResponse(
        path=jd_path,
        media_type="application/pdf",
        filename=download_name,
    )


@router.post(
    "/{drive_id}/round-results",
)
async def upload_round_results(
    drive_id: int,
    stage: str = Form(...),
    file: UploadFile = File(...),
    current_user: User = Depends(
        require_admin
    ),
    db: Session = Depends(
        get_db
    ),
):
    drive = (
        db.query(
            PlacementDrive
        )
        .filter(
            PlacementDrive.id
            == drive_id
        )
        .first()
    )

    if not drive:
        raise HTTPException(
            status_code=404,
            detail=(
                "Placement drive not found"
            ),
        )

    stage = normalize_stage(stage)

    allowed_stages = list(ROUND_STAGES)

    if stage not in allowed_stages:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid recruitment round. "
                "Allowed rounds are: "
                + ", ".join(
                    allowed_stages
                )
            ),
        )

    if (
        stage == "Resume Shortlisting"
        and not drive.resume_shortlisting
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Resume shortlisting is "
                "not enabled for this "
                "placement drive"
            ),
        )

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected",
        )

    try:
        validate_extension_and_mime(
            file,
            EXCEL_MIME_TYPES,
            (
                "Only .xlsx and .xlsm "
                "Excel files are allowed"
            ),
        )

        contents = await read_upload_with_limit(
            file,
            MAX_EXCEL_SIZE,
        )

        if not contents:
            raise HTTPException(
                status_code=400,
                detail=(
                    "The uploaded Excel "
                    "file is empty"
                ),
            )

        validate_excel_content(
            contents
        )

        workbook = load_workbook(
            filename=BytesIO(contents),
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        workbook.close()

    except HTTPException:
        raise

    except Exception as error:
        logger.exception(
            "Unable to parse uploaded round-results file"
        )

        raise HTTPException(
            status_code=400,
            detail="Unable to read the uploaded Excel file.",
        ) from error

    finally:
        await file.close()

    if not rows:
        raise HTTPException(
            status_code=400,
            detail=(
                "The Excel file "
                "contains no data"
            ),
        )

    headers = []

    for value in rows[0]:
        if value is None:
            headers.append("")
        else:
            headers.append(
                str(value)
                .strip()
                .lower()
            )

    roll_no_index = None

    for index, header in enumerate(
        headers
    ):
        normalized_header = (
            header
            .replace("_", " ")
            .replace("-", " ")
            .replace(".", "")
            .strip()
        )

        normalized_header = (
            " ".join(
                normalized_header.split()
            )
        )

        if normalized_header in {
            "roll no",
            "roll number",
            "rollno",
            "roll",
        }:
            roll_no_index = index
            break

    if roll_no_index is None:
        raise HTTPException(
            status_code=400,
            detail=(
                "Excel file must contain "
                "a Roll No. column"
            ),
        )

    excel_roll_numbers = set()

    for row in rows[1:]:

        if (
            roll_no_index
            >= len(row)
        ):
            continue

        value = row[
            roll_no_index
        ]

        if value is None:
            continue

        roll_no = normalize_roll_no(
            value
        )

        if roll_no:
            excel_roll_numbers.add(
                roll_no
            )

    if not excel_roll_numbers:
        raise HTTPException(
            status_code=400,
            detail=(
                "No student roll numbers "
                "were found in the Excel file"
            ),
        )

    next_stage = next_stage_after_round(
        stage
    )

    applications = (
        db.query(
            Application,
            Student,
        )
        .join(
            Student,
            Student.id
            == Application.student_id,
        )
        .filter(
            Application.drive_id
            == drive_id,
        )
        .all()
    )

    current_stage_applications = []

    for (
        application,
        student,
    ) in applications:

        if not is_application_ready_for_round(
            application,
            drive.resume_shortlisting,
            stage,
        ):
            continue

        current_stage_applications.append(
            (
                application,
                student,
            )
        )

    if not current_stage_applications:

        available_stages = sorted(
            {
                normalize_stage(
                    application.current_stage
                )
                for (
                    application,
                    student,
                ) in applications
                if (
                    application.status
                    not in {
                        "Rejected",
                        "Selected",
                    }
                )
            }
        )

        if available_stages:

            raise HTTPException(
                status_code=400,
                detail=(
                    f"No applications are "
                    f"currently ready for "
                    f"the '{stage}' round. "
                    f"Current application "
                    f"stages are: "
                    f"{', '.join(available_stages)}"
                ),
            )

        raise HTTPException(
            status_code=400,
            detail=(
                "No applications exist "
                "for this placement drive"
            ),
        )

    passed_count = 0
    failed_count = 0

    processed_students = []

    for (
        application,
        student,
    ) in current_stage_applications:

        student_roll_no = (
            normalize_roll_no(
                student.roll_no
            )
        )

        if (
            student_roll_no
            in excel_roll_numbers
        ):

            passed_count += 1

            advance_application_for_round(
                application,
                drive.resume_shortlisting,
                stage,
            )

            processed_students.append(
                {
                    "roll_no":
                        student.roll_no,
                    "result":
                        "passed",
                }
            )

        else:

            failed_count += 1

            reject_application_for_round(
                application,
                drive.resume_shortlisting,
                stage,
            )

            processed_students.append(
                {
                    "roll_no":
                        student.roll_no,
                    "result":
                        "rejected",
                }
            )

    try:

        db.commit()

    except Exception as error:

        db.rollback()

        logger.exception(
            "Unexpected error while processing round results"
        )

        raise HTTPException(
            status_code=500,
            detail="Failed to process round results.",
        ) from error

    return {
        "message": (
            "Round results processed "
            "successfully"
        ),
        "drive_id":
            drive_id,
        "stage":
            stage,
        "next_stage":
            next_stage,
        "passed_count":
            passed_count,
        "failed_count":
            failed_count,
        "total_processed":
            (
                passed_count
                + failed_count
            ),
        "uploaded_filename":
            file.filename,
        "processed_students":
            processed_students,
    }
